#!/usr/bin/env python3
"""Generate Ansible inventory variables from network-config.xlsx.

The Excel workbook is the editable source of truth. This script validates the
workbook before writing generated YAML files under inventory/.
"""

from __future__ import annotations

import argparse
import ipaddress
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "network-config.xlsx"
INVENTORY_DIR = PROJECT_ROOT / "inventory"
GENERATED_HEADER = "# GENERATED FILE - edit network-config.xlsx instead.\n"


class ConfigError(ValueError):
    """Raised when workbook configuration is invalid."""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate network-config.xlsx and generate Ansible YAML."
    )
    parser.add_argument(
        "--file",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to the Excel workbook (default: network-config.xlsx)",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="Validate only; do not write YAML files.",
    )
    return parser.parse_args()


def is_enabled(value: Any) -> bool:
    if value is None or value == "":
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return str(value).strip().lower() in {"true", "yes", "1", "enabled", "on"}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_table(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ConfigError(f"Missing required sheet: {sheet_name}")

    sheet = workbook[sheet_name]
    headers = [clean(cell.value) for cell in sheet[1]]
    if not any(headers):
        raise ConfigError(f"Sheet {sheet_name} has no headers")

    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        row = dict(zip(headers, values))
        row["_row"] = excel_row
        rows.append(row)
    return rows


def validate_ip(value: Any, label: str) -> str:
    text = clean(value)
    try:
        return str(ipaddress.ip_address(text))
    except ValueError as exc:
        raise ConfigError(f"{label}: invalid IP address '{text}'") from exc


def validate_network(value: Any, label: str) -> ipaddress.IPv4Network:
    text = clean(value)
    try:
        network = ipaddress.ip_network(text, strict=True)
    except ValueError as exc:
        raise ConfigError(f"{label}: invalid network '{text}'") from exc
    if network.version != 4:
        raise ConfigError(f"{label}: only IPv4 is supported")
    return network


def load_config(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise ConfigError(f"Workbook not found: {path}")

    workbook = load_workbook(path, data_only=True)

    settings_rows = read_table(workbook, "Settings")
    settings = {
        clean(row.get("Key")): row.get("Value")
        for row in settings_rows
        if clean(row.get("Key"))
    }

    required_settings = {
        "domain_name",
        "dns_server",
        "snmp_community",
        "syslog_host",
        "syslog_level",
        "dhcp_server_device",
    }
    missing_settings = sorted(required_settings - settings.keys())
    if missing_settings:
        raise ConfigError(
            "Settings sheet is missing: " + ", ".join(missing_settings)
        )

    validate_ip(settings["dns_server"], "Settings/dns_server")
    validate_ip(settings["syslog_host"], "Settings/syslog_host")

    # Devices
    devices: dict[str, dict[str, Any]] = {}
    for row in read_table(workbook, "Devices"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"Devices row {row['_row']}"
        name = clean(row.get("Device"))
        device_type = clean(row.get("Type"))
        if not name:
            raise ConfigError(f"{line}: Device is required")
        if name in devices:
            raise ConfigError(f"{line}: duplicate device '{name}'")
        if device_type not in {"router", "core_switch", "access_switch"}:
            raise ConfigError(f"{line}: invalid Type '{device_type}'")
        devices[name] = {
            "type": device_type,
            "ansible_host": validate_ip(row.get("Management IP"), f"{line}/Management IP"),
            "create_svis": is_enabled(row.get("Create SVIs")),
        }

    if not devices:
        raise ConfigError("No enabled devices found")

    dhcp_server_device = clean(settings["dhcp_server_device"])
    if dhcp_server_device not in devices:
        raise ConfigError(
            f"Settings/dhcp_server_device references unknown device '{dhcp_server_device}'"
        )

    # VLANs
    vlans: list[dict[str, Any]] = []
    vlan_ids: set[int] = set()
    for row in read_table(workbook, "VLANs"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"VLANs row {row['_row']}"
        try:
            vlan_id = int(row.get("VLAN ID"))
        except (TypeError, ValueError) as exc:
            raise ConfigError(f"{line}: VLAN ID must be an integer") from exc
        if not 1 <= vlan_id <= 4094:
            raise ConfigError(f"{line}: VLAN ID {vlan_id} is outside 1-4094")
        if vlan_id in vlan_ids:
            raise ConfigError(f"{line}: duplicate VLAN ID {vlan_id}")
        vlan_ids.add(vlan_id)

        name = clean(row.get("Name"))
        if not name:
            raise ConfigError(f"{line}: Name is required")

        network = validate_network(row.get("Subnet"), f"{line}/Subnet")
        gateway = ipaddress.ip_address(validate_ip(row.get("Gateway"), f"{line}/Gateway"))
        if gateway not in network or gateway in {network.network_address, network.broadcast_address}:
            raise ConfigError(f"{line}: gateway {gateway} is not a usable host in {network}")

        dhcp_enabled = is_enabled(row.get("DHCP Enabled"))
        excluded_start = clean(row.get("DHCP Exclude Start"))
        excluded_end = clean(row.get("DHCP Exclude End"))
        if dhcp_enabled:
            if not excluded_start or not excluded_end:
                raise ConfigError(f"{line}: DHCP exclusion start/end are required")
            start_ip = ipaddress.ip_address(validate_ip(excluded_start, f"{line}/DHCP Exclude Start"))
            end_ip = ipaddress.ip_address(validate_ip(excluded_end, f"{line}/DHCP Exclude End"))
            if start_ip not in network or end_ip not in network or start_ip > end_ip:
                raise ConfigError(f"{line}: invalid DHCP exclusion range")

        vlans.append(
            {
                "id": vlan_id,
                "name": name,
                "subnet": str(network),
                "gateway": str(gateway),
                "create_svi": is_enabled(row.get("Create SVI")),
                "dhcp_enabled": dhcp_enabled,
                "dhcp_excluded_start": excluded_start,
                "dhcp_excluded_end": excluded_end,
            }
        )

    if not vlans:
        raise ConfigError("No enabled VLANs found")

    # Ports
    ports_by_device: dict[str, dict[str, list[dict[str, Any]]]] = defaultdict(
        lambda: {"trunk_interfaces": [], "access_interfaces": []}
    )
    seen_interfaces: set[tuple[str, str]] = set()
    for row in read_table(workbook, "Ports"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"Ports row {row['_row']}"
        device = clean(row.get("Device"))
        interface = clean(row.get("Interface"))
        mode = clean(row.get("Mode")).lower()
        if device not in devices:
            raise ConfigError(f"{line}: unknown device '{device}'")
        if devices[device]["type"] == "router":
            raise ConfigError(f"{line}: router '{device}' cannot use switch port definitions")
        if not interface:
            raise ConfigError(f"{line}: Interface is required")
        key = (device, interface)
        if key in seen_interfaces:
            raise ConfigError(f"{line}: duplicate interface {device}/{interface}")
        seen_interfaces.add(key)

        description = clean(row.get("Description"))
        if mode == "access":
            try:
                access_vlan = int(row.get("Access VLAN"))
            except (TypeError, ValueError) as exc:
                raise ConfigError(f"{line}: Access VLAN must be an integer") from exc
            if access_vlan not in vlan_ids:
                raise ConfigError(f"{line}: Access VLAN {access_vlan} does not exist")
            ports_by_device[device]["access_interfaces"].append(
                {"name": interface, "vlan": access_vlan, "description": description}
            )
        elif mode == "trunk":
            allowed_text = clean(row.get("Allowed VLANs"))
            try:
                allowed = [int(part.strip()) for part in allowed_text.split(",") if part.strip()]
            except ValueError as exc:
                raise ConfigError(f"{line}: Allowed VLANs must be comma-separated integers") from exc
            if not allowed:
                raise ConfigError(f"{line}: Allowed VLANs is required for trunk mode")
            unknown = sorted(set(allowed) - vlan_ids)
            if unknown:
                raise ConfigError(f"{line}: unknown allowed VLAN(s): {unknown}")
            ports_by_device[device]["trunk_interfaces"].append(
                {"name": interface, "allowed_vlans": ",".join(str(v) for v in allowed)}
            )
        else:
            raise ConfigError(f"{line}: Mode must be access or trunk")

    # Routes
    routes_by_device: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in read_table(workbook, "Routes"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"Routes row {row['_row']}"
        device = clean(row.get("Device"))
        if device not in devices:
            raise ConfigError(f"{line}: unknown device '{device}'")
        prefix = validate_ip(row.get("Prefix"), f"{line}/Prefix")
        mask = validate_ip(row.get("Mask"), f"{line}/Mask")
        try:
            ipaddress.ip_network(f"{prefix}/{mask}", strict=False)
        except ValueError as exc:
            raise ConfigError(f"{line}: invalid prefix/mask combination") from exc
        next_hop = validate_ip(row.get("Next Hop"), f"{line}/Next Hop")
        routes_by_device[device].append(
            {"prefix": prefix, "mask": mask, "next_hop": next_hop}
        )

    # ACLs
    acl_rows: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in read_table(workbook, "ACLs"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"ACLs row {row['_row']}"
        device = clean(row.get("Device"))
        name = clean(row.get("ACL Name"))
        action = clean(row.get("Action")).lower()
        rule = clean(row.get("Rule"))
        apply_to = clean(row.get("Apply To"))
        direction = clean(row.get("Direction")).lower()
        if device not in devices:
            raise ConfigError(f"{line}: unknown device '{device}'")
        if not name or not rule or not apply_to:
            raise ConfigError(f"{line}: ACL Name, Rule and Apply To are required")
        if action not in {"permit", "deny"}:
            raise ConfigError(f"{line}: Action must be permit or deny")
        if direction not in {"in", "out"}:
            raise ConfigError(f"{line}: Direction must be in or out")
        try:
            sequence = int(row.get("Sequence"))
        except (TypeError, ValueError) as exc:
            raise ConfigError(f"{line}: Sequence must be an integer") from exc
        acl_rows[(device, name)].append(
            {
                "sequence": sequence,
                "rule": f"{action} {rule}",
                "apply_to": apply_to,
                "direction": direction,
            }
        )

    acls_by_device: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (device, name), entries in acl_rows.items():
        apply_targets = {(entry["apply_to"], entry["direction"]) for entry in entries}
        if len(apply_targets) != 1:
            raise ConfigError(f"ACL '{name}' on {device} has inconsistent Apply To/Direction")
        apply_to, direction = next(iter(apply_targets))
        acls_by_device[device].append(
            {
                "name": name,
                "rules": [entry["rule"] for entry in sorted(entries, key=lambda item: item["sequence"])],
                "apply_to": apply_to,
                "direction": direction,
            }
        )

    # NAT
    nat_by_device: dict[str, dict[str, Any]] = {}
    for row in read_table(workbook, "NAT"):
        if not is_enabled(row.get("Enabled")):
            continue
        line = f"NAT row {row['_row']}"
        device = clean(row.get("Device"))
        if device not in devices:
            raise ConfigError(f"{line}: unknown device '{device}'")
        acl_name = clean(row.get("ACL Name"))
        network = clean(row.get("Network"))
        outside = clean(row.get("Outside Interface"))
        if not acl_name or not network or not outside:
            raise ConfigError(f"{line}: ACL Name, Network and Outside Interface are required")
        current = nat_by_device.setdefault(
            device,
            {"nat_acl_name": acl_name, "nat_networks": [], "nat_outside_interface": outside},
        )
        if current["nat_acl_name"] != acl_name or current["nat_outside_interface"] != outside:
            raise ConfigError(f"{line}: NAT rows for {device} must use one ACL and outside interface")
        current["nat_networks"].append(network)

    return {
        "settings": settings,
        "devices": devices,
        "vlans": vlans,
        "ports_by_device": ports_by_device,
        "routes_by_device": routes_by_device,
        "acls_by_device": acls_by_device,
        "nat_by_device": nat_by_device,
        "dhcp_server_device": dhcp_server_device,
    }


def dump_yaml(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    yaml_text = yaml.safe_dump(
        data,
        sort_keys=False,
        allow_unicode=True,
        default_flow_style=False,
        width=100,
    )
    path.write_text(GENERATED_HEADER + "---\n" + yaml_text, encoding="utf-8")


def build_hosts(devices: dict[str, dict[str, Any]]) -> dict[str, Any]:
    routers: dict[str, Any] = {}
    core_switches: dict[str, Any] = {}
    access_switches: dict[str, Any] = {}

    for name, device in devices.items():
        entry = {"ansible_host": device["ansible_host"]}
        if device["type"] == "router":
            routers[name] = entry
        elif device["type"] == "core_switch":
            core_switches[name] = entry
        else:
            access_switches[name] = entry

    return {
        "all": {
            "children": {
                "routers": {"hosts": routers},
                "switches": {
                    "children": {
                        "core_switches": {"hosts": core_switches},
                        "access_switches": {"hosts": access_switches},
                    }
                },
            }
        }
    }


def generate(config: dict[str, Any]) -> list[Path]:
    written: list[Path] = []
    settings = config["settings"]

    network_data = {
        "domain_name": clean(settings["domain_name"]),
        "dns_server": clean(settings["dns_server"]),
        "snmp_community": clean(settings["snmp_community"]),
        "syslog_host": clean(settings["syslog_host"]),
        "syslog_level": clean(settings["syslog_level"]),
        "vlans": [
            {
                "id": vlan["id"],
                "name": vlan["name"],
                "subnet": vlan["subnet"],
                "gateway": vlan["gateway"],
                "create_svi": vlan["create_svi"],
            }
            for vlan in config["vlans"]
        ],
    }
    network_path = INVENTORY_DIR / "group_vars" / "all" / "network.yml"
    dump_yaml(network_path, network_data)
    written.append(network_path)

    hosts_path = INVENTORY_DIR / "hosts.yml"
    dump_yaml(hosts_path, build_hosts(config["devices"]))
    written.append(hosts_path)

    for device_name, device in config["devices"].items():
        host_data: dict[str, Any] = {}

        if device["type"] != "router":
            host_data["create_svis"] = device["create_svis"]

        ports = config["ports_by_device"].get(device_name, {})
        trunks = ports.get("trunk_interfaces", [])
        access = ports.get("access_interfaces", [])
        if trunks:
            host_data["trunk_interfaces"] = trunks
        if access:
            host_data["access_interfaces"] = access

        routes = config["routes_by_device"].get(device_name, [])
        if routes:
            host_data["static_routes"] = routes

        acls = config["acls_by_device"].get(device_name, [])
        if acls:
            host_data["acls"] = acls

        if device_name == config["dhcp_server_device"]:
            dhcp_vlans = [vlan for vlan in config["vlans"] if vlan["dhcp_enabled"]]
            if dhcp_vlans:
                host_data["dhcp_excluded"] = [
                    {"start": vlan["dhcp_excluded_start"], "end": vlan["dhcp_excluded_end"]}
                    for vlan in dhcp_vlans
                ]
                host_data["dhcp_pools"] = []
                for vlan in dhcp_vlans:
                    network = ipaddress.ip_network(vlan["subnet"])
                    host_data["dhcp_pools"].append(
                        {
                            "name": f"VLAN{vlan['id']}",
                            "network": str(network.network_address),
                            "mask": str(network.netmask),
                            "gateway": vlan["gateway"],
                        }
                    )

        nat = config["nat_by_device"].get(device_name)
        if nat:
            host_data.update(nat)

        host_path = INVENTORY_DIR / "host_vars" / f"{device_name}.yml"
        dump_yaml(host_path, host_data)
        written.append(host_path)

    return written


def main() -> int:
    args = parse_args()
    try:
        config = load_config(args.file.resolve())
        print(f"Validation OK: {args.file}")
        print(f"  Devices: {len(config['devices'])}")
        print(f"  VLANs:   {len(config['vlans'])}")

        if args.check:
            print("Check-only mode: no files were changed.")
            return 0

        written = generate(config)
        print("Generated:")
        for path in written:
            print(f"  - {path.relative_to(PROJECT_ROOT)}")
        return 0
    except ConfigError as exc:
        print(f"ERROR: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
